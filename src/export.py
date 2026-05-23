"""Export Excel du screening — palette finance + tableau Tufte dense."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
from openpyxl.utils import get_column_letter

# Palette du skill dashboard-finance
BRAND = "0B2545"
INK_900 = "0A0E1A"
INK_700 = "2A3142"
INK_300 = "9BA3B5"
INK_100 = "E5E8EE"
ACCENT = "C8102E"
POSITIVE = "0F7A3A"


def _format_dirigeants_pp(dirigeants: list[dict]) -> str:
    pp = [d for d in dirigeants if d.get("type_dirigeant") == "personne physique"
          and "commissaire" not in (d.get("qualite") or "").lower()]
    parts = []
    for d in pp:
        prenoms = (d.get("prenoms") or "").split(",")[0].strip()
        parts.append(f"{prenoms} {d.get('nom', '')} ({d.get('qualite', '')})".strip())
    return "; ".join(parts)


def _format_dirigeants_pm(dirigeants: list[dict]) -> str:
    pm = [d for d in dirigeants if d.get("type_dirigeant") == "personne morale"
          and "commissaire" not in (d.get("qualite") or "").lower()]
    return "; ".join(f"{d.get('denomination', '')} ({d.get('qualite', '')})" for d in pm)


def entreprise_to_row(e: dict) -> dict:
    siege = e.get("siege") or {}
    dirigeants = e.get("dirigeants") or []
    siren = e.get("siren")
    return {
        "SIREN": siren,
        "Dénomination": e.get("nom_complet") or e.get("nom_raison_sociale"),
        "Forme": e.get("nature_juridique"),
        "NAF": e.get("activite_principale"),
        "Effectif (tranche)": e.get("tranche_effectif_salarie"),
        "Réf. effectif": e.get("annee_tranche_effectif_salarie"),
        "Date création": (e.get("date_creation") or "")[:10],
        "Ancienneté (ans)": e.get("_anciennete"),
        "Adresse": siege.get("adresse"),
        "CP": siege.get("code_postal"),
        "Commune": siege.get("libelle_commune"),
        "Département": siege.get("departement"),
        "Région": siege.get("region"),
        "Dirigeants PP (direction)": _format_dirigeants_pp(dirigeants),
        "Personnes morales liées (direction)": _format_dirigeants_pm(dirigeants),
        "Diagnostic indépendance": e.get("_independance_motif"),
        "Motif rejet": e.get("_rejet_motif"),
        "Annuaire (à vérifier âge dirigeant)": (
            f"https://annuaire-entreprises.data.gouv.fr/dirigeants/{siren}" if siren else None
        ),
    }


def ecrire_excel(
    retenues: list[dict],
    rejetees: list[dict],
    parametres: dict,
    output_path: Path,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    retenues_tri = sorted(retenues, key=lambda e: e.get("_anciennete", 0), reverse=True)
    df_ret = pd.DataFrame([entreprise_to_row(e) for e in retenues_tri])
    df_rej = pd.DataFrame([entreprise_to_row(e) for e in rejetees])

    total = len(retenues) + len(rejetees)
    taux = (len(retenues) / total * 100) if total else 0
    synthese = pd.DataFrame([
        ("Date du screening", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("NAF ciblé", parametres.get("naf")),
        ("Périmètre", parametres.get("geo")),
        ("Tranche effectif", f"{parametres.get('effectif_min')} - {parametres.get('effectif_max')}"),
        ("Ancienneté minimale", f"{parametres.get('anciennete_min')} ans (proxy âge dirigeant)"),
        ("Indépendance exigée", "oui" if parametres.get("exclure_filiales_groupe") else "non"),
        ("Population brute", total),
        ("Cibles retenues tranche 1", len(retenues)),
        ("Cibles rejetées", len(rejetees)),
        ("Taux de retenue", f"{taux:.1f} %"),
        ("", ""),
        ("À FAIRE — étape suivante",
         "Vérifier manuellement l'âge du dirigeant via la colonne 'Annuaire' "
         "(la date de naissance n'est pas exposée par l'API publique recherche-entreprises ; "
         "échantillon 2 = enrichissement via API INPI RNE avec token)."),
    ], columns=["Indicateur", "Valeur"])

    motifs = (
        df_rej["Motif rejet"].fillna("?")
        .str.split(":").str[0].str.split("(").str[0]
        .str.strip().value_counts()
        .rename_axis("Motif").reset_index(name="Nombre")
        if not df_rej.empty else pd.DataFrame(columns=["Motif", "Nombre"])
    )

    with pd.ExcelWriter(output_path, engine="openpyxl") as w:
        synthese.to_excel(w, sheet_name="Synthèse", index=False)
        (df_ret.drop(columns=["Motif rejet"], errors="ignore")
               .to_excel(w, sheet_name="Cibles retenues", index=False))
        motifs.to_excel(w, sheet_name="Motifs de rejet", index=False)
        df_rej.to_excel(w, sheet_name="Cibles rejetées", index=False)
        pd.DataFrame(list(parametres.items()), columns=["Paramètre", "Valeur"]).to_excel(
            w, sheet_name="Paramètres", index=False)

    _polish(output_path)
    return output_path


def _polish(path: Path) -> None:
    wb = load_workbook(path)
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color=BRAND, end_color=BRAND, fill_type="solid")
    cell_font = Font(name="Calibri", size=10)
    thin = Side(border_style="thin", color=INK_100)

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = cell_font
                cell.alignment = Alignment(vertical="top")
        for col in ws.columns:
            values = [c.value for c in col if c.value is not None]
            length = max((len(str(v)) for v in values), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(length + 2, 12), 70)
        ws.row_dimensions[1].height = 28
        ws.freeze_panes = "A2"

    wb.save(path)
